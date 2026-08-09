"""Market-wide corporate calendar — results, dividends, and other actions for
ALL listed stocks (not just the watchlist), exported to Excel.

Reuses the same NSE feeds the events poller reads, but here we do NOT filter to
the watchlist and we keep the calendar-relevant records (a date + a purpose)
across every symbol NSE returns.
"""
from __future__ import annotations

import io
from datetime import datetime

from .nse import NSEClient
from . import events

# purpose -> normalised action type, for the corp_action feed
_ACTION_TYPES = [
    ("dividend", "Dividend"), ("bonus", "Bonus"),
    ("split", "Split"), ("sub-division", "Split"), ("subdivision", "Split"),
    ("buyback", "Buyback"), ("buy-back", "Buyback"),
    ("rights", "Rights"), ("amalgamation", "Scheme"), ("demerger", "Scheme"),
    ("capital reduction", "Scheme"),
]


def _action_type(text: str) -> str:
    t = text.lower()
    for kw, label in _ACTION_TYPES:
        if kw in t:
            return label
    return "Other"


def _company(rec: dict) -> str:
    return events._first(rec, "sm_name", "comp", "companyName", "company", default="")


def build(client: NSEClient | None = None) -> list[dict]:
    """Return calendar rows: {symbol, company, type, date, detail, source}."""
    nse = client or NSEClient()
    rows: list[dict] = []

    # board meetings -> results / other meeting purposes, dated
    try:
        for rec in nse.fetch_feed("board_meeting"):
            n = events.normalize("board_meeting", rec)
            if not n or not n["when"]:
                continue
            purpose = n["headline"]
            typ = "Results" if any(k in purpose.lower() for k in
                  ("financial result", "results", "unaudited", "audited")) else "Board meeting"
            rows.append({"symbol": n["symbol"], "company": _company(rec),
                         "type": typ, "date": n["when"], "detail": purpose,
                         "source": "Board meeting"})
    except Exception as e:  # keep whatever we got from the other feed
        rows.append({"symbol": "", "company": "", "type": "ERROR",
                     "date": "", "detail": f"board_meeting feed failed: {e}",
                     "source": ""})

    # corporate actions -> dividends / bonus / split / buyback ..., with ex-date
    try:
        for rec in nse.fetch_feed("corp_action"):
            n = events.normalize("corp_action", rec)
            if not n:
                continue
            subject = n["headline"]
            ex = events._first(rec, "exDate", "exdate")
            rec_date = events._first(rec, "recDate", "recordDate")
            date = ex or rec_date or n["when"]
            if not date:
                continue
            detail = subject
            if rec_date:
                detail = f"{subject}  (record {rec_date})"
            rows.append({"symbol": n["symbol"], "company": _company(rec),
                         "type": _action_type(subject), "date": date,
                         "detail": detail, "source": "Corporate action"})
    except Exception as e:
        rows.append({"symbol": "", "company": "", "type": "ERROR",
                     "date": "", "detail": f"corp_action feed failed: {e}",
                     "source": ""})

    rows.sort(key=lambda r: (_sortkey(r["date"]), r["symbol"]))
    return rows


def _sortkey(d: str):
    """Best-effort parse of NSE date strings for sorting; unknowns go last."""
    for fmt in ("%d-%b-%Y", "%d-%m-%Y", "%Y-%m-%d", "%d %b %Y",
                "%d-%b-%Y %H:%M:%S"):
        try:
            return (0, datetime.strptime(d.strip()[:len(fmt)+4], fmt))
        except (ValueError, AttributeError):
            continue
    return (1, datetime.max)


def to_xlsx(rows: list[dict]) -> bytes:
    """Render calendar rows to an .xlsx workbook (bytes)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Corporate calendar"
    headers = ["Date", "Symbol", "Company", "Type", "Details", "Source"]
    ws.append(headers)

    head_fill = PatternFill("solid", fgColor="1C2A41")
    head_font = Font(bold=True, color="E6EDF3")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill; cell.font = head_font
        cell.alignment = Alignment(vertical="center")

    for r in rows:
        ws.append([r["date"], r["symbol"], r["company"], r["type"],
                   r["detail"], r["source"]])

    widths = [16, 14, 34, 14, 60, 18]
    for i, w in enumerate(widths, 1):
        ws.column_letter = get_column_letter(i)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
